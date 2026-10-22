"""Tien ich python-docx dung de sinh bao cao (tieu de, doan van, bang, hinh, code)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Times New Roman"
BODY_PT = 14
FIRST_LINE_CM = 1.27  # thut dau dong doan van than bai
TEXT_WIDTH_CM = 16.0  # A4 21 cm - le trai 3 cm - le phai 2 cm


class Report:
    def __init__(self):
        self.doc = Document()
        self.fig_no = 0
        self.tab_no = 0
        self.headings: list[tuple[int, str]] = []  # (cap, tieu de) theo thu tu xuat hien -> muc luc
        self._setup_styles()

    # ---- style ---------------------------------------------------------------------------
    def _setup_styles(self):
        d = self.doc
        for s in d.sections:
            s.page_width, s.page_height = Cm(21.0), Cm(29.7)  # A4 (mac dinh python-docx la Letter)
            s.top_margin = s.bottom_margin = Cm(2.0)
            s.left_margin = Cm(3.0)
            s.right_margin = Cm(2.0)
        # Kich thuoc, gian dong theo mau bao cao cua khoa (TNR 14, H1 20 can giua, H2 16, H3 14, chu thich 13)
        cap = d.styles["Caption"]
        cap.font.name = FONT
        cap.font.size = Pt(13)
        cap.font.bold = False
        cap.font.color.rgb = RGBColor(0, 0, 0)
        cap.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        cap.paragraph_format.space_before = Pt(4)
        cap.paragraph_format.space_after = Pt(12)
        cap.paragraph_format.line_spacing = 1.0
        st = d.styles["Normal"]
        st.font.name = FONT
        st.font.size = Pt(BODY_PT)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.15
        for lvl, size, before, after in ((1, 20, 0, 12), (2, 16, 18, 6), (3, 14, 12, 5)):
            h = d.styles[f"Heading {lvl}"]
            h.font.name = FONT
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = False
            h.font.color.rgb = RGBColor(0, 0, 0)
            h.font.all_caps = lvl == 1
            if lvl == 1:
                h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fonts = h.element.rPr.rFonts
            fonts.set(qn("w:eastAsia"), FONT)
            for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):  # theme font de len FONT
                fonts.attrib.pop(qn(attr), None)
            h.paragraph_format.space_before = Pt(before)
            h.paragraph_format.space_after = Pt(after)
        for name in ("List Bullet", "List Number"):
            d.styles[name].font.name = FONT
            d.styles[name].font.size = Pt(BODY_PT)

    # ---- khoi noi dung ---------------------------------------------------------------------
    def h(self, level: int, text: str):
        """Tieu de; cap 1 luon bat dau trang moi (khong chen them neu vua co ngat trang)."""
        if level == 1 and not self._at_page_start():
            self.page_break()
        self.headings.append((level, text))
        return self.doc.add_heading(text, level=level)

    def _at_page_start(self) -> bool:
        body = self.doc.element.body
        last = body[-2] if len(body) > 1 else None  # phan tu cuoi la sectPr
        return last is None or bool(last.xpath('.//w:br[@w:type="page"]'))

    def cover_frame(self, path: str | Path, width_cm: float, height_cm: float):
        """Anh khung bia nam sau chu, can giua vung le (giong mau bia PTIT)."""
        para = self.doc.add_paragraph()
        para.paragraph_format.space_after = Pt(0)
        inline = para.add_run().add_picture(str(path), width=Cm(width_cm), height=Cm(height_cm))._inline
        anchor = OxmlElement("wp:anchor")
        for k, v in (("distT", "0"), ("distB", "0"), ("distL", "0"), ("distR", "0"), ("simplePos", "0"),
                     ("relativeHeight", "0"), ("behindDoc", "1"), ("locked", "0"), ("layoutInCell", "1"),
                     ("allowOverlap", "1")):
            anchor.set(k, v)
        pos = OxmlElement("wp:simplePos")
        pos.set("x", "0")
        pos.set("y", "0")
        anchor.append(pos)
        for tag in ("wp:positionH", "wp:positionV"):
            el = OxmlElement(tag)
            el.set("relativeFrom", "margin")
            al = OxmlElement("wp:align")
            al.text = "center"
            el.append(al)
            anchor.append(el)
        anchor.append(inline.extent)
        anchor.append(OxmlElement("wp:wrapNone"))
        for child in (inline.docPr, inline.find(qn("wp:cNvGraphicFramePr")), inline.graphic):
            if child is not None:
                anchor.append(child)
        inline.getparent().replace(inline, anchor)
        return para

    def toc_line(self, text: str, page: str, level: int = 1):
        """Mot dong muc luc: tieu de ... so trang (tab phai co dau cham dan)."""
        para = self.doc.add_paragraph()
        fmt = para.paragraph_format
        fmt.left_indent = Cm(0.75 * (level - 1))
        fmt.space_after = Pt(2)
        fmt.tab_stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = para.add_run(f"{text}\t{page}")
        run.bold = run.font.all_caps = level == 1
        return para

    def page_numbers(self):
        """So trang o giua chan trang, an tren trang bia."""
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        para = sec.footer.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run()
        for tag, attr in (("w:fldChar", "begin"), ("w:instrText", None), ("w:fldChar", "separate"),
                          ("w:t", None), ("w:fldChar", "end")):
            el = OxmlElement(tag)
            if tag == "w:fldChar":
                el.set(qn("w:fldCharType"), attr)
            elif tag == "w:instrText":
                el.set(qn("xml:space"), "preserve")
                el.text = " PAGE "
            else:
                el.text = "1"
            run._r.append(el)

    def p(self, text: str = "", bold=False, italic=False, align=None, size=None, style=None):
        para = self.doc.add_paragraph(style=style)
        if text:
            run = para.add_run(text)
            run.bold, run.italic = bold, italic
            if size:
                run.font.size = Pt(size)
        if align == "center":
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "left":
            para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if text and style is None and not bold:  # doan than bai: thut dong dau
                para.paragraph_format.first_line_indent = Cm(FIRST_LINE_CM)
        return para

    def rich(self, parts: list[tuple[str, dict]]):
        """parts = [(text, {"bold":..,"italic":..,"code":..}), ...]"""
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.first_line_indent = Cm(FIRST_LINE_CM)
        for text, fmt in parts:
            run = para.add_run(text)
            run.bold = fmt.get("bold", False)
            run.italic = fmt.get("italic", False)
            if fmt.get("code"):
                run.font.name = "Consolas"
                run.font.size = Pt(11)
        return para

    def bullets(self, items: list[str], style="List Bullet"):
        for it in items:
            para = self.doc.add_paragraph(style=style)
            para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if isinstance(it, tuple):
                r = para.add_run(it[0])
                r.bold = True
                para.add_run(it[1])
            else:
                para.add_run(it)

    def numbered(self, items: list[str]):
        """Moi danh sach danh so lai tu 1 (mac dinh Word noi tiep so giua cac danh sach cung style)."""
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        style_num = self.doc.styles["List Number"].element.pPr.numPr.numId.val
        num = numbering.add_num(numbering.num_having_numId(style_num).abstractNumId.val)
        num.add_lvlOverride(ilvl=0).add_startOverride(1)
        self.bullets(items, style="List Number")
        for para in self.doc.paragraphs[-len(items):]:
            numpr = para._p.get_or_add_pPr().get_or_add_numPr()
            numpr.get_or_add_ilvl().val = 0
            numpr.get_or_add_numId().val = num.numId

    def code(self, text: str, size=9.5):
        for line in text.rstrip("\n").split("\n"):
            para = self.doc.add_paragraph()
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            run = para.add_run(line if line else " ")
            run.font.name = "Consolas"
            run.font.size = Pt(size)
            self._shade(para, "F2F2F2")
        self.doc.add_paragraph()

    @staticmethod
    def _shade(para, hex_color: str):
        pPr = para._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        pPr.append(shd)

    def table(self, headers: list[str], rows: list[list], widths: list[float] | None = None, caption: str = "",
              size=10.5, align_center_cols: set[int] | None = None):
        if caption:
            self.tab_no += 1
            cap = self.doc.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(f"Bảng {self.tab_no}. {caption}")
            r.bold = r.italic = True
        t = self.doc.add_table(rows=1, cols=len(headers))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, h in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(h))
            run.bold = True
            run.font.size = Pt(size)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            self._shade(cell.paragraphs[0], "DCE6F1")
            self._cell_shade(cell, "DCE6F1")
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                cells[i].text = ""
                run = cells[i].paragraphs[0].add_run(str(v))
                run.font.size = Pt(size)
                if align_center_cols and i in align_center_cols:
                    cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if widths:
            t.autofit = False
            for col, w in zip(t._tbl.tblGrid.gridCol_lst, widths):  # LibreOffice doc do rong tu tblGrid
                col.w = Cm(w)
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Cm(w)
        self.doc.add_paragraph()
        return t

    @staticmethod
    def _cell_shade(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hex_color)
        tcPr.append(shd)

    def figure(self, path: str | Path, caption: str, width_cm: float = 16.0):
        path = Path(path)
        if not path.exists():
            self.p(f"[thiếu hình: {path.name}]", italic=True)
            return
        self.fig_no += 1
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(path), width=Cm(width_cm))
        cap = self.doc.add_paragraph(style="Caption")
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Hình {self.fig_no}. {caption}")
        r.italic = True

    def page_break(self):
        self.doc.add_page_break()

    def save(self, path: str | Path):
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return Path(path)
